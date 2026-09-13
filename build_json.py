#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
公告 XLSX -> announcement.json

用法:
    python build_json.py            读取 XLSX，生成 announcement.json
    python build_json.py --check    只打印将要生成的 JSON，不写文件
    python build_json.py --fix      修复 XLSX 里被破坏的韩文（如果出现 ?）
    python build_json.py --init     重建两份 XLSX 模板

数据源（都在本目录）:
    DT_公告数据表.xlsx              四国语言公告
                                    列：编号,详情英文,详情中文,详情日文,详情韩文
    DT_更新奖励数据表.xlsx          这次更新发的奖励
                                    列按 ST_日常任务奖励结构体：
                                    ---,奖励名称,ID,奖励数量,物品图标,奖励?,
                                    说明,中文说明,已领取,活跃值,日文说明,韩文说明
    version_code.txt                这次的 versionCode（纯数字，如 185）

为什么用 XLSX（而不是 CSV）:
    CSV 在 Excel 里默认存成 GBK/ANSI，而**谚文不在 GBK 字符集**，一存就变成 ? 且无法恢复。
    XLSX 是 Unicode 原生存储（Office Open XML），中文/日文/韩文都不会丢，彻底根治乱码。
    需要 openpyxl：pip install openpyxl
"""

import os
import sys
import json
import re
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.stderr.write(
        "缺少依赖 openpyxl。请先运行：\n"
        "  pip install openpyxl\n"
        "（本机 venv 已装：%APPDATA%\\..\\.workbuddy\\binaries\\python\\envs\\default\\Scripts\\python.exe -m pip install openpyxl）\n"
    )
    sys.exit(2)

HERE = os.path.dirname(os.path.abspath(__file__))
NOTICE_XLSX = os.path.join(HERE, "DT_公告数据表.xlsx")
REWARD_XLSX = os.path.join(HERE, "DT_更新奖励数据表.xlsx")
CODE_TXT = os.path.join(HERE, "version_code.txt")
JSON_OUT = os.path.join(HERE, "announcement.json")

# 与 ST_公告结构体 字段一一对应（顺序按工程内 DataTable 的字段顺序）
NOTICE_LANGS = ["详情英文", "详情中文", "详情日文", "详情韩文"]
# 输出 JSON 时的顺序（中文在前方便阅读；蓝图按字段名取值，顺序无影响）
OUT_ORDER = ["详情中文", "详情英文", "详情日文", "详情韩文"]

# 韩文模板（与工程内 DT_公告数据表 的原文一致，版本号动态替换）
KO_TEMPLATE = (
    "버전:{ver}\n"
    "1.스킬 피해 범위 강화；\n"
    "2.자원 획득 용이；\n"
    "3.광고 응답 속도 향상；\n"
    "4.여러 버그 수정。"
)

SAMPLE_NOTICE = {
    "编号": "1",
    "详情英文": "Version:V 0.1.6.3\n1.Enhance skill damage range;\n2.Make resources easier to obtain;\n3.Improve ad response speed;\n4.Fix several bugs.",
    "详情中文": "版本号:V 0.1.6.3\n1.增强技能伤害范围；\n2.资源获得更容易；\n3.提升广告响应速度：\n4.修复若干Bug。",
    "详情日文": "バージョン:V 0.1.6.3\n1.スキルのダメージ範囲を強化；\n2.リソースの獲得を容易に；\n3.広告の応答速度を向上；\n4.いくつかのバグを修正。",
    "详情韩文": KO_TEMPLATE.format(ver="V 0.1.6.3"),
}

# 奖励表列：与 ST_日常任务奖励结构体 完全一致
REWARD_FIELDS = [
    "---", "奖励名称", "ID", "奖励数量", "物品图标", "奖励?",
    "说明", "中文说明", "已领取", "活跃值", "日文说明", "韩文说明",
]

_NS = '[8F3A4F7F78E52C175B3305EF4AA5E433]'
_ICON = "/Script/Engine.Texture2D'/Game/BreakTheBricks/Art/UI/主页/Texture/"


def _loc(key, val):
    return 'NSLOCTEXT("' + _NS + '", "' + key + '", "' + val + '")'


SAMPLE_REWARDS = [
    {
        "---": "1",
        "奖励名称": "金币",
        "ID": "3",
        "奖励数量": "1000.000000",
        "物品图标": _ICON + "T_UI_Item_Gtid_金币_01a.T_UI_Item_Gtid_金币_01a'",
        "奖励?": "False",
        "说明": _loc("E35E5A4345B6FDC8C025618594E31B76", "Gold"),
        "中文说明": _loc("DA80D0DD429DDC149D37B591EBEEEBB8", "金币"),
        "已领取": "False",
        "活跃值": "0.000000",
        "日文说明": _loc("E089AACF4A1D47D04315DFBB39E3AF4B", "ゴールド"),
        "韩文说明": _loc("2D0324114E3A67D5C24350A2CA930399", "골드"),
    },
    {
        "---": "2",
        "奖励名称": "钻石",
        "ID": "4",
        "奖励数量": "600.000000",
        "物品图标": _ICON + "T_UI_Item_Gtid_钻石_01b.T_UI_Item_Gtid_钻石_01b'",
        "奖励?": "False",
        "说明": _loc("1AD715994D385EA2CB6D94B3D92DFF8A", "Diamonds"),
        "中文说明": _loc("E5C20F38442990207213DFA25E9324C7", "钻石"),
        "已领取": "False",
        "活跃值": "0.000000",
        "日文说明": _loc("7B65EAA447EB125C669D5DBCB3AB1FDA", "ダイヤ"),
        "韩文说明": _loc("1F28F7774025D33FDEC4AA8652809EB6", "다이아"),
    },
]


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\r\n", "\n").replace("\r", "\n").strip()


def _int(v, default=0):
    try:
        if v is None or str(v).strip() == "":
            return default
        return int(float(str(v)))
    except (TypeError, ValueError):
        return default


def read_xlsx(path):
    """读 XLSX 第一张表，返回 [dict, ...]（表头->单元格）。找不到返回 None。

    注意：刻意不用 read_only=True。当表被 Excel 撑大到很大的 used range
    （例如 A1:T200，20 列 × 200 行，多为空单元格）时，openpyxl 的 read_only
    解析会只读到首行首列，导致数据“凭空消失”、rewards 静默变空。
    我们这份文件很小，普通模式即可，且更稳。
    """
    if not os.path.exists(path):
        return None
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None:
            continue
        # 跳过整行全空的（Excel 撑大的空行），避免无意义的 None 行
        if all((c is None or (isinstance(c, str) and c.strip() == "")) for c in r):
            continue
        d = {}
        for i, h in enumerate(headers):
            d[h] = r[i] if i < len(r) else None
        out.append(d)
    return out


def write_xlsx(path, fields, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(fields)
    for r in rows:
        ws.append([r.get(f, "") for f in fields])
    wb.save(path)


def pick_version_name(cn_text):
    """从中文正文第一行「版本号:V x.y.z」里取版本名"""
    for line in cn_text.splitlines():
        line = line.strip()
        for prefix in ("版本号:", "版本号："):
            if line.startswith(prefix):
                return line[len(prefix):].strip()
    return ""


# 四语言版本行前缀。正文首行形如「版本号:V 0.1.6.4」「Version: V 0.1.6.4」等。
VER_LINE_PREFIXES = ("版本号:", "版本号：", "Version:", "Version：", "バージョン:", "버전:")
_VER_NUM_RE = re.compile(r"\d+(?:\.\d+)+")


def is_version_line(line):
    s = line.strip()
    return any(s.startswith(p) for p in VER_LINE_PREFIXES)


def set_notice_version(langs, ver):
    """把四语正文首个版本行的版本号改写为 ver（保留原有前缀与「V 」风格）。
    例：「版本号:V 0.1.6.4」+ 0.1.6.5 → 「版本号:V 0.1.6.5」。
    只改第一个版本行的数字部分，其余行（含描述内容）一律不动。"""
    if not (ver or "").strip():
        return langs
    ver = ver.strip()
    out = dict(langs)
    for k, v in out.items():
        if not isinstance(v, str):
            continue
        lines = v.replace("\r\n", "\n").split("\n")
        for i, line in enumerate(lines):
            if is_version_line(line):
                m = _VER_NUM_RE.search(line)
                if m:
                    lines[i] = line[:m.start()] + ver + line[m.end():]
                    out[k] = "\n".join(lines)
                break
    return out


def fix_notice_xlsx():
    rows = read_xlsx(NOTICE_XLSX)
    if rows is None:
        print("找不到:", NOTICE_XLSX)
        return 1
    fixed = 0
    for r in rows:
        cn = _text(r.get("详情中文"))
        ko = _text(r.get("详情韩文"))
        if "?" in ko:
            r["详情韩文"] = KO_TEMPLATE.format(ver=pick_version_name(cn))
            fixed += 1
    write_xlsx(NOTICE_XLSX, ["编号"] + NOTICE_LANGS, rows)
    print("已重写:", NOTICE_XLSX)
    print("  恢复被破坏的韩文行数:", fixed)
    return 0


def fix_reward_xlsx():
    rows = read_xlsx(REWARD_XLSX)
    if rows is None:
        print("找不到:", REWARD_XLSX)
        return 0
    sample_by_id = {s["ID"]: s for s in SAMPLE_REWARDS}
    fixed = 0
    for r in rows:
        s = sample_by_id.get(_text(r.get("ID")))
        if not s:
            continue
        for k in ("韩文说明", "日文说明"):
            v = _text(r.get(k))
            if v and "?" in v:
                r[k] = s[k]
                fixed += 1
    write_xlsx(REWARD_XLSX, REWARD_FIELDS, rows)
    print("已重写:", REWARD_XLSX)
    print("  恢复字段数:", fixed)
    return 0


def write_template():
    if not os.path.exists(NOTICE_XLSX):
        write_xlsx(NOTICE_XLSX, ["编号"] + NOTICE_LANGS, [SAMPLE_NOTICE])
        print("已生成:", NOTICE_XLSX)
    else:
        print("已存在，跳过:", NOTICE_XLSX)

    if not os.path.exists(REWARD_XLSX):
        write_xlsx(REWARD_XLSX, REWARD_FIELDS, SAMPLE_REWARDS)
        print("已生成:", REWARD_XLSX)
    else:
        print("已存在，跳过:", REWARD_XLSX)

    if not os.path.exists(CODE_TXT):
        with open(CODE_TXT, "w", encoding="utf-8", newline="") as f:
            f.write("185")
        print("已生成:", CODE_TXT)
    return 0


def read_version_code():
    if os.path.exists(CODE_TXT):
        with open(CODE_TXT, "r", encoding="utf-8-sig") as f:
            return _int(_text(f.read()), 0)
    if os.path.exists(JSON_OUT):
        try:
            with open(JSON_OUT, "r", encoding="utf-8") as f:
                return _int(json.load(f).get("version_code"), 0)
        except Exception:
            return 0
    return 0


def build(override_version_name=None, override_version_code=None):
    n_rows = read_xlsx(NOTICE_XLSX)
    if n_rows is None:
        print("找不到公告表：%s" % NOTICE_XLSX)
        print("先运行：python build_json.py --init")
        return None
    if not n_rows:
        print("公告表是空的：%s" % NOTICE_XLSX)
        return None

    n_fields = list(n_rows[0].keys())
    missing = [c for c in NOTICE_LANGS if c not in n_fields]
    if missing:
        print("公告表缺少列：%s" % ", ".join(missing))
        print("当前列：%s" % ", ".join(n_fields))
        return None

    warns = []

    n = n_rows[0]
    langs = {k: _text(n.get(k)) for k in NOTICE_LANGS}
    if "?" in langs["详情韩文"]:
        # 极少数情况下 XLSX 里仍然出现 ?（比如曾被当成 CSV 存过），先用内置正确韩文兜底
        langs["详情韩文"] = KO_TEMPLATE.format(ver=pick_version_name(langs["详情中文"]))
        warns.append(
            "详情韩文 里出现 ?，已用内置正确韩文兜底。请确认 DT_公告数据表.xlsx 是用 Excel 另存为 .xlsx（不是 CSV）。"
        )

    notice = {"Name": _text(n.get("编号")) or "1"}
    # 打包流水线传入游戏版本（--version-name）时，同步改写四语正文的版本行：
    #   勾选递增 → 写"本次将打出的版本"；取消递增 → 写当前 UE 配置版本。
    # 版本行以外的描述内容仍以数据源 xlsx 为准，不做任何改动。
    if (override_version_name or "").strip():
        langs = set_notice_version(langs, override_version_name.strip())
    for k in OUT_ORDER:
        notice[k] = langs[k]

    rewards = []
    r_rows = read_xlsx(REWARD_XLSX)
    if r_rows is None:
        warns.append(
            "找不到奖励表 DT_更新奖励数据表.xlsx（应与本脚本同目录），rewards 将为空。"
            "请确认文件没被改名/移动。"
        )
    elif not r_rows:
        warns.append(
            "奖励表 DT_更新奖励数据表.xlsx 读不到任何有效数据行，rewards 将为空。"
            "若表里明明有数据，多半是表被 Excel 撑得过大导致解析异常——本脚本已改用普通模式读取，"
            "重新运行一次即可；如仍为空，请检查首列/ID 列是否为空。"
        )
    else:
        for r in r_rows:
            if not _text(r.get("ID")):
                continue
            rewards.append({
                "ID": _text(r.get("ID")),
                "奖励名称": _text(r.get("奖励名称")),
                "奖励数量": _int(r.get("奖励数量"), 0),
            })

    # 版本号：默认从公告文本 `版本号:V x.y.z` / version_code.txt 提取；
    # 若打包流水线传入了游戏版本（--version-name / --version-code）则优先采用，实现"自动匹配"。
    # 注意：只匹配版本号，描述内容（notice / rewards）始终以数据源 xlsx 为准，这里不做任何改动。
    vname = (override_version_name or "").strip() or pick_version_name(langs["详情中文"])
    vcode = override_version_code if override_version_code is not None else read_version_code()

    data = {
        "schema": 1,
        "version_name": vname,
        "version_code": vcode,
        "updated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "notice": notice,
        "rewards": rewards,
    }
    if warns:
        data["_warning"] = warns
    return data


def dump_tables():
    """把两个 XLSX 的全部数据按行打印出来，方便人工核对（不写文件）。

    用法：python build_json.py --dump
    """
    targets = [
        ("公告数据表", NOTICE_XLSX),
        ("更新奖励数据表", REWARD_XLSX),
    ]
    for label, path in targets:
        print("=" * 64)
        print("【%s】  %s" % (label, os.path.basename(path)))
        rows = read_xlsx(path)
        if rows is None:
            print("  !! 文件不存在（应与本脚本同目录）")
            continue
        if not rows:
            print("  (空表，没有任何数据行)")
            continue
        headers = list(rows[0].keys())
        for idx, r in enumerate(rows, 1):
            print("  ---- 第 %d 行 ----" % idx)
            for h in headers:
                v = r.get(h)
                if v is None or (isinstance(v, str) and v.strip() == ""):
                    continue
                s = str(v).replace("\r\n", " / ").replace("\n", " / ").replace("\r", " / ")
                print("      %s = %s" % (h, s))
    print("=" * 64)


def main():
    args = sys.argv[1:]

    if "--init" in args:
        return write_template()
    if "--fix" in args:
        fix_notice_xlsx()
        return fix_reward_xlsx()
    if "--dump" in args:
        dump_tables()
        return 0

    # 打包流水线可传入游戏版本，让公告版本号自动对齐（只匹配版本号，描述内容不变）
    vname_arg = None
    vcode_arg = None
    if "--version-name" in args:
        i = args.index("--version-name")
        if i + 1 < len(args):
            vname_arg = args[i + 1]
    if "--version-code" in args:
        i = args.index("--version-code")
        if i + 1 < len(args):
            vcode_arg = _int(args[i + 1])

    data = build(vname_arg, vcode_arg)
    if data is None:
        return 1

    if "--check" in args:
        print(json.dumps(data, ensure_ascii=False, indent=2))
        return 0

    with open(JSON_OUT, "w", encoding="utf-8", newline="\n") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    # 回写 version_code.txt，保持文件与本次生成的版本号一致（仅覆盖数字，不动描述内容）
    if vcode_arg is not None:
        try:
            with open(CODE_TXT, "w", encoding="utf-8", newline="") as f:
                f.write(str(data["version_code"]))
        except Exception:
            pass
    matched = "游戏版本自动匹配" if (vname_arg or vcode_arg is not None) else "公告文本/version_code.txt"
    print("已生成：%s" % JSON_OUT)
    print("  版本：%s (versionCode %s)  [来源: %s]" % (data["version_name"], data["version_code"], matched))
    for k in OUT_ORDER:
        print("  %s：%d 字" % (k, len(data["notice"][k])))
    for r in data["rewards"]:
        print("  奖励：%s x %d (ID=%s)" % (r["奖励名称"], r["奖励数量"], r["ID"]))
    for w in data.get("_warning", []):
        print("  [警告] %s" % w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
